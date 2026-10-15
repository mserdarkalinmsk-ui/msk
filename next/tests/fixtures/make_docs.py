"""Synthetic project documents (all content invented) for document-reading tests."""
import docx, openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table
from reportlab.lib.styles import getSampleStyleSheet

d = docx.Document()
d.add_heading('Description of the Action', 0)
d.add_paragraph('Title of the action: "Demo Rights Programme for Testing"')
d.add_paragraph('The implementation period of the action shall be 24 months and shall start on 1 March 2026.')
d.add_heading('Action plan', 1)
t = d.add_table(rows=1, cols=8); h = ['Activity', 'Lead'] + [f'M{i}' for i in range(1, 7)]
for c, v in zip(t.rows[0].cells, h): c.text = v
rows = [['Output 1: Laws reviewed', '', '', '', '', '', '', ''],
        ['1.1 Needs assessment', 'CoE', 'P', 'P', 'X', 'X', '', ''],
        ['1.2 Expert review', 'CoE', '', '', 'X', 'X', 'X', ''],
        ['Output 2: Capacity built', '', '', '', '', '', '', ''],
        ['2.1 Training of trainers', 'CoE', '', '', '', 'X', 'X', 'X']]
for r in rows:
    cells = t.add_row().cells
    for c, v in zip(cells, r): c.text = v
d.add_heading('Risks', 1)
t = d.add_table(rows=1, cols=4)
for c, v in zip(t.rows[0].cells, ['Risk', 'Likelihood', 'Impact', 'Mitigation']): c.text = v
for r in [['Staff turnover', 'High', 'Medium', 'Backup staffing'], ['Political change', 'Low', 'Very high', 'Dialogue']]:
    for c, v in zip(t.add_row().cells, r): c.text = v
d.add_heading('Logical framework', 1)
t = d.add_table(rows=1, cols=4)
for c, v in zip(t.rows[0].cells, ['Result', 'Indicator', 'Baseline', 'Target']): c.text = v
for r in [['Outcome 1', 'Number of laws aligned', '0', '3'], ['', 'Judges trained', '0', '200']]:
    for c, v in zip(t.add_row().cells, r): c.text = v
d.save('tests/fixtures/doa.docx')

doc = SimpleDocTemplate('tests/fixtures/contract.pdf', pagesize=A4); st = getSampleStyleSheet()
doc.build([
    Paragraph('CONTRIBUTION AGREEMENT', st['Title']),
    Paragraph('Contract No 700008888', st['Normal']),
    Paragraph('The total estimated eligible costs of the action are EUR 2 500 000. The maximum contribution of the European Union is EUR 2 250 000, which is 90 % of the total estimated eligible costs.', st['Normal']),
    Paragraph('The implementation period of the action shall be 24 months. The implementation shall start on 01/03/2026.', st['Normal']),
    Paragraph('The organisation shall submit a progress report every 12 months, within 60 days after the end of each reporting period. Pre-financing of 30 % is paid on signature.', st['Normal']),
    Table([['Risk', 'Likelihood', 'Impact'], ['Delays in procurement', '3', '4']]),
])

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Action plan'
ws.append(['Workplan']); ws.append(['Code', 'Activity'] + ['Mar-26', 'Apr-26', 'May-26', 'Jun-26'])
ws.append(['Result 1']); ws.append(['1.1', 'Kick-off', 'x', '', '', '']); ws.append(['1.2', 'Survey', '', 'x', 'x', ''])
wb.save('tests/fixtures/plan.xlsx')
